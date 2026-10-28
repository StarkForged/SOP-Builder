"""Parse uploaded notes and transcript files into text or speaker turns."""

import io
import re

import filetype
import webvtt
from docx import Document
from pypdf import PdfReader

TEXT_EXT = {".txt", ".md", ".vtt", ".srt"}
ALLOWED_EXT = TEXT_EXT | {".docx", ".pdf"}
SPEAKER = re.compile(r"^\s*([A-Z][A-Za-z0-9 .'_-]{0,30}?)\s*:\s+(.+)$")


class BadFile(ValueError):
    pass


def ext_of(name: str) -> str:
    m = re.search(r"\.[A-Za-z0-9]+$", name or "")
    return m.group(0).lower() if m else ""


def check_magic(name: str, data: bytes) -> str:
    ext = ext_of(name)
    if ext not in ALLOWED_EXT:
        raise BadFile(f"Unsupported file type '{ext or name}'. Use .txt .md .docx .pdf .vtt .srt.")
    if ext == ".pdf" and not data.startswith(b"%PDF"):
        raise BadFile("File content is not a PDF.")
    if ext == ".docx" and not data.startswith(b"PK\x03\x04"):
        raise BadFile("File content is not a Word .docx document.")
    if ext in TEXT_EXT and (b"\x00" in data[:4096] or filetype.guess(data) is not None):
        raise BadFile("File content is not plain text.")
    return ext


def decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def extract_text(name: str, data: bytes) -> str:
    ext = check_magic(name, data)
    try:
        if ext == ".docx":
            doc = Document(io.BytesIO(data))
            lines = []
            for p in doc.paragraphs:
                style = (p.style.name or "") if p.style is not None else ""
                lines.append(("- " if "List" in style else "") + p.text)
            for table in doc.tables:
                for row in table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(lines)
        if ext == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:  # corrupt document
        raise BadFile("Could not read the document; it may be corrupt.") from exc
    return decode_text(data)


def _split_speaker(text: str, voice: str | None = None) -> tuple[str | None, str]:
    text = re.sub(r"\s+", " ", text).strip()
    if voice:
        return voice.strip(), text
    m = SPEAKER.match(text)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return None, text


def _secs(ts: str) -> float:
    parts = [float(x) for x in ts.replace(",", ".").split(":")]
    total = 0.0
    for p in parts:
        total = total * 60 + p
    return round(total, 3)


def _merge(turns: list[dict]) -> list[dict]:
    out: list[dict] = []
    for t in turns:
        if not t["text"]:
            continue
        if out and t["speaker"] and out[-1]["speaker"] == t["speaker"]:
            out[-1]["text"] += " " + t["text"]
            out[-1]["end"] = t.get("end")
        else:
            out.append(dict(t))
    return out


def parse_transcript(name: str, data: bytes) -> list[dict]:
    """Return speaker turns: {speaker, text, start, end}."""
    ext = check_magic(name, data)
    if ext in (".vtt", ".srt"):
        text = decode_text(data).replace("\r\n", "\n")
        try:
            captions = webvtt.from_buffer(io.StringIO(text), format="srt" if ext == ".srt" else "vtt")
        except Exception as exc:
            raise BadFile(f"Could not parse the {ext} transcript.") from exc
        turns = []
        for c in captions:
            speaker, body = _split_speaker(c.text, getattr(c, "voice", None))
            turns.append({"speaker": speaker, "text": body,
                          "start": _secs(c.start), "end": _secs(c.end)})
        return _merge(turns)
    return parse_transcript_text(extract_text(name, data))


def parse_transcript_text(text: str) -> list[dict]:
    turns: list[dict] = []
    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            continue
        speaker, body = _split_speaker(line)
        if speaker is None and turns and turns[-1]["speaker"]:
            turns[-1]["text"] += " " + body
            continue
        turns.append({"speaker": speaker, "text": body, "start": None, "end": None})
    return _merge(turns)
