import io

from docx import Document

from app.redact import redact
from app.segments import role_from_question, split_note_lines


def make_pdf(lines):
    content = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({x}) Tj T*" for x in lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def test_notes_segmenting_and_wrapped_lines(client):
    text = ("Refund process\n\n- Open the Orders page and search for the\n  order number\n"
            "- Click Issue refund,\nthen confirm\n2. Check the status\n\n* Email the customer")
    lines = split_note_lines(text)
    assert lines == ["Refund process", "Open the Orders page and search for the order number",
                     "Click Issue refund, then confirm", "Check the status", "Email the customer"]
    r = client.post("/api/notes", data={"source_index": "3", "text": text})
    assert r.status_code == 200
    segs = r.json()["segments"]
    assert [s["id"] for s in segs][:2] == ["S3-0001", "S3-0002"]
    assert all(s["kind"] == "note_line" and s["speaker_role"] == "staff" for s in segs)


def test_docx_pdf_vtt_srt_parsing(client):
    doc = Document()
    doc.add_paragraph("Open the dashboard")
    doc.add_paragraph("Click Export", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    r = client.post("/api/notes", data={"source_index": "1"},
                    files={"file": ("notes.docx", buf.getvalue(), "application/octet-stream")})
    assert [s["text"] for s in r.json()["segments"]] == ["Open the dashboard", "Click Export"]

    r = client.post("/api/notes", data={"source_index": "2"},
                    files={"file": ("n.pdf", make_pdf(["Open the ledger", "Post the entry"]), "application/pdf")})
    assert r.status_code == 200
    texts = " ".join(s["text"] for s in r.json()["segments"])
    assert "Open the ledger" in texts and "Post the entry" in texts

    vtt = ("WEBVTT\n\n00:00:01.000 --> 00:00:04.000\nAsha: How do you start?\n\n"
           "00:00:04.500 --> 00:00:09.000\nRavi: I open the dashboard.\n\n"
           "00:00:09.000 --> 00:00:12.000\nRavi: Then I click Export.\n")
    r = client.post("/api/notes", data={"source_index": "4", "kind": "transcript"},
                    files={"file": ("i.vtt", vtt.encode(), "text/vtt")})
    body = r.json()
    segs = body["segments"]
    assert len(segs) == 2 and segs[1]["text"] == "I open the dashboard. Then I click Export."
    assert segs[1]["start_s"] == 4.5 and segs[1]["end_s"] == 12.0
    assert body["speakers"] == {"S4-0001": "Asha", "S4-0002": "Ravi"}

    srt = ("1\n00:00:01,000 --> 00:00:03,000\nAsha: What next?\n\n"
           "2\n00:00:03,000 --> 00:00:06,500\nRavi: Submit the form.\n")
    r = client.post("/api/notes", data={"source_index": "5", "kind": "transcript"},
                    files={"file": ("i.srt", srt.encode(), "text/plain")})
    segs = r.json()["segments"]
    assert segs[1]["text"] == "Submit the form." and segs[1]["start_s"] == 3.0
    assert r.json()["speakers"]["S5-0002"] == "Ravi"

    r = client.post("/api/notes", data={"source_index": "6", "kind": "transcript",
                                        "text": "Asha: Who approves?\nRavi: The manager\napproves it."})
    segs = r.json()["segments"]
    assert segs[1]["text"] == "The manager approves it."


def test_reject_oversize_and_bad_magic_bytes(client):
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 100
    for name in ("notes.pdf", "notes.txt", "notes.docx"):
        r = client.post("/api/notes", data={"source_index": "1"}, files={"file": (name, png, "x/y")})
        assert r.status_code == 400 and r.json()["error"] == "bad_file"
    r = client.post("/api/notes", data={"source_index": "1"}, files={"file": ("a.exe", b"MZ", "x/y")})
    assert r.status_code == 400
    big = b"a" * 4_000_001
    r = client.post("/api/notes", data={"source_index": "1"}, files={"file": ("a.txt", big, "text/plain")})
    assert r.status_code == 400 and "too large" in r.json()["message"]
    r = client.post("/api/notes", data={"source_index": "1"},
                    files={"file": ("a.txt", b"a" * 4_500_000, "text/plain")})
    assert r.status_code == 413
    r = client.post("/api/notes", data={"source_index": "1"},
                    files={"file": ("broken.pdf", b"%PDF-1.4 garbage", "application/pdf")})
    assert r.status_code == 400
    r = client.post("/api/notes", data={"source_index": "1", "text": "   "})
    assert r.status_code == 400


def test_redaction_types_and_money_kept():
    tok = "gsk" + "_" + "A1b2C3d4E5f6G7h8J9k0"
    text = (f"Mail ravi.k@example.com or call +91 98765 43210. Card 4111 1111 1111 1111. "
            f"Aadhaar 2345 6789 0123. PAN ABCDE1234F. password: hunter2 key {tok}. "
            "Refunds above ₹5,000 or Rs 1,00,000 or $2500.50 or INR 12000000 need approval. 15% fee.")
    out = redact(text)
    for bad in ("ravi.k@example.com", "98765", "4111", "2345 6789", "ABCDE1234F", "hunter2", tok):
        assert bad not in out
    for tag in ("[EMAIL]", "[PHONE]", "[CARD]", "[AADHAAR]", "[PAN]", "[SECRET]"):
        assert tag in out
    for kept in ("₹5,000", "Rs 1,00,000", "$2500.50", "INR 12000000", "15%"):
        assert kept in out
    assert "1234 5678 9012 3456" in redact("Order ref 1234 5678 9012 3456")  # fails Luhn -> not a card


def test_redaction_idempotent():
    text = "Email a@b.co, phone 9876543210, token=abc123xyz, pwd: x, Rs 5,000"
    once = redact(text)
    assert redact(once) == once
    assert redact("") == ""


def test_interviewer_tag_from_question():
    assert role_from_question("How do you approve a refund?") == "interviewer"
    assert role_from_question("I open the dashboard.") == "staff"


def test_static_index_served(client):
    r = client.get("/")
    assert r.status_code == 200 and "<html" in r.text.lower()
    assert client.get("/api/health").json()["ok"] is True


def test_sample_endpoint_and_js_mime(client):
    r = client.get("/api/sample")
    assert r.status_code == 200 and r.json()["title"] and "approval" in r.json()["text"]
    r = client.get("/app.js")
    assert r.status_code == 200 and "javascript" in r.headers["content-type"]
    assert "css" in client.get("/style.css").headers["content-type"]


def test_fake_ai_mode_runs_sample_end_to_end(client, monkeypatch):
    monkeypatch.setenv("FAKE_AI", "1")
    sample = client.get("/api/sample").json()
    segs = client.post("/api/notes", data={"source_index": "1", "text": sample["text"]}).json()["segments"]
    r = client.post("/api/generate", json={"title_hint": sample["title"], "segments": segs})
    assert r.status_code == 200, r.text
    body = r.json()
    assert len(body["sop"]["steps"]) >= 4 and body["sop"]["warnings"]
    cats = {q["category"] for q in body["sop"]["questions"]}
    assert {"ambiguous_term", "permission_access", "failure_handling"} <= cats
    assert body["report"]["model"] == "demo (FAKE_AI)"
    from tests.test_media import make_wav

    r = client.post("/api/transcribe", data={"source_index": "2", "consent": "true"},
                    files={"chunk": ("c.wav", make_wav(1), "audio/wav")})
    assert len(r.json()["segments"]) == 2
    r = client.post("/api/frames", json={"source_index": 3, "items": [
        {"text": "Orders > Issue refund", "t": 0, "confidence": 90}]})
    assert r.json()["segments"][0]["kind"] == "screen_text"
